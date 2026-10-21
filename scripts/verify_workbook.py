#!/usr/bin/env python3
"""
verify_workbook.py - Step 6 of the website-clicks-improver skill.

Recalculates the workbook with LibreOffice (if installed) and checks:
  * every expected tab exists
  * no formula errors (#REF!, #VALUE!, #DIV/0!, #NAME?, #N/A)
  * suggested titles <= 60 chars and metas between 70 and 160 chars
  * no "(not written yet)" chunks left (warning only)

Usage:  python verify_workbook.py outputs/report.xlsx
Exit code 1 if a hard check fails.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

TABS = ["Summary", "Low CTR Pages", "Page 2-3 Pages", "Keyword Gaps", "Content Chunks", "Internal Links",
        "Action Plan", "Method & Notes"]
ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "Err:")


def recalc(path: Path) -> tuple[Path | None, str]:
    """Returns (recalculated copy, status) where status is 'ok', 'missing' or 'failed'."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None, "missing"
    out = Path(tempfile.mkdtemp())
    profile = Path(tempfile.mkdtemp())  # private profile: parallel runs can't fight over a lock
    try:
        res = subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--calc",
                              "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                             capture_output=True, timeout=240)
    except subprocess.TimeoutExpired:
        return None, "failed"
    f = out / path.name
    return (f, "ok") if res.returncode == 0 and f.exists() else (None, "failed")


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = Path(sys.argv[1])
    hard, soft = [], []

    wb = load_workbook(path)
    for t in TABS:
        if t not in wb.sheetnames:
            hard.append(f"Missing tab: {t}")

    ws = wb["Low CTR Pages"] if "Low CTR Pages" in wb.sheetnames else None
    if ws:
        for r in range(2, ws.max_row + 1):
            url = ws[f"A{r}"].value
            # Q = title, S = alt title, U = meta, W = alt meta
            for col, kind, required in (("Q", "title", True), ("S", "alternative title", False),
                                        ("U", "meta", True), ("W", "alternative meta", False)):
                v = str(ws[f"{col}{r}"].value or "")
                if not v:
                    if required:
                        soft.append(f"No suggested {kind}: {url}")
                elif "title" in kind and len(v) > 60:
                    hard.append(f"Suggested {kind} is {len(v)} chars (>60): {url}")
                elif "meta" in kind and not 70 <= len(v) <= 160:
                    hard.append(f"Suggested {kind} is {len(v)} chars (want 70-160): {url}")
    if "Content Chunks" in wb.sheetnames:
        ws = wb["Content Chunks"]
        empty = sum(1 for r in range(2, ws.max_row + 1) if ws[f"I{r}"].value == "(not written yet)")
        if empty:
            soft.append(f"{empty} content chunk(s) not written yet")

    calc, status = recalc(path)
    if status == "failed":
        hard.append("LibreOffice is installed but recalculation failed - formulas were NOT checked. Re-run verify.")
    if calc:
        wbv = load_workbook(calc, data_only=True)
        n = 0
        for sh in wbv.worksheets:
            for row in sh.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith(ERRORS):
                        hard.append(f"Formula error {c.value} at {sh.title}!{c.coordinate}")
                    n += 1
        s = wbv["Summary"]
        print("Summary KPIs after recalculation:")
        for r in range(4, 15):
            if s[f"A{r}"].value:
                print(f"  {s[f'A{r}'].value:<42} {s[f'B{r}'].value}")
    elif status == "missing":
        soft.append("LibreOffice not found - formulas not recalculated (they will calculate when opened in Excel/Sheets)")

    for m in soft:
        print("WARN :", m)
    for m in hard:
        print("FAIL :", m)
    print("RESULT:", "FAIL" if hard else "PASS")
    sys.exit(1 if hard else 0)


if __name__ == "__main__":
    main()
