#!/usr/bin/env python3
"""
build_report.py - Step 5 of the website-clicks-improver skill.

Builds the client-ready Excel working sheet from:
  analysis.json     (analyze_gsc.py)
  pages.json        (fetch_pages.py)
  gaps.json         (find_gaps.py)
  suggestions.json  (written by Claude: titles, metas, diagnoses, push plans, content chunks)

Tabs: Summary | Low CTR Pages | Page 2-3 Pages | Keyword Gaps | Content Chunks | Internal Links |
      Cannibalisation (if any) | Action Plan | Method & Notes

suggestions.json shape (every key optional):
{
  "site_name": "Acme Plumbing",
  "pages": {
    "<url>": {
      "diagnosis": "why CTR is low / why the page is stuck",
      "title_options": ["recommended title", "alternative"],
      "meta_options": ["recommended meta", "alternative"],
      "h1_suggestion": "...",
      "other_fixes": ["..."],
      "actions": ["push-to-page-1 step", "..."]
    }
  },
  "chunks": [
    {"url": "<url>", "cluster": "C2", "keywords": ["..."], "chunk_type": "...", "placement": "...",
     "heading": "...", "content": "...", "notes": "facts to confirm"}
  ]
}

Usage:
  python build_report.py --analysis work/analysis.json --pages work/pages.json --gaps work/gaps.json
         --suggestions work/suggestions.json -o outputs/clicks-improver.xlsx [--md outputs/summary.md]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, INK, MUTED = "0D1730", "1F2937", "6B7280"
CYAN, VIOLET, GREEN, AMBER, RED = "22D3EE", "A78BFA", "34D399", "FBBF24", "F87171"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
SUGGEST_FILL = PatternFill("solid", fgColor="ECFDF5")   # Claude's suggestions = light green
INPUT_FILL = PatternFill("solid", fgColor="FFF7E6")     # columns for the user to fill = light amber
KPI_FILL = PatternFill("solid", fgColor="F1F5F9")
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
PRIORITY_FILL = {"High": "FEE2E2", "Medium": "FEF3C7", "Low": "E0F2FE"}
STATUSES = '"To do,In progress,Done,Skipped"'
ATTRIBUTION = "Built with website-clicks-improver by Akif Qureshi | SEO Specialist - algoblueprints.com"


def short(url: str) -> str:
    p = urlparse(url)
    return (p.path or "/") + (("?" + p.query) if p.query else "")


def table(ws, headers: list[tuple[str, int]], rows: list[list], start_row: int = 1,
          suggest_cols: tuple = (), input_cols: tuple = (), status_col: int | None = None) -> int:
    """Write a styled table. headers = [(name, width)]. Returns last row number."""
    for ci, (name, width) in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=ci, value=name)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[start_row].height = 32
    for ri, row in enumerate(rows, start_row + 1):
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            if isinstance(val, str) and val.startswith("="):
                c.data_type = "s"  # text that happens to start with '=' must not become a formula
            c.alignment, c.border = WRAP, BORDER
            c.font = Font(size=10, color=INK)
            if ci in suggest_cols:
                c.fill = SUGGEST_FILL
            elif ci in input_cols:
                c.fill = INPUT_FILL
    last = start_row + len(rows)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{last}"
    if status_col and rows:
        dv = DataValidation(type="list", formula1=STATUSES, allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(status_col)
        dv.add(f"{col}{start_row + 1}:{col}{max(last, start_row + 1)}")
        ws.conditional_formatting.add(f"{col}{start_row + 1}:{col}{last}",
                                      CellIsRule(operator="equal", formula=['"Done"'],
                                                 fill=PatternFill("solid", fgColor="D1FAE5")))
    return last


def fmt(ws, col: str, first: int, last: int, number_format: str) -> None:
    for r in range(first, last + 1):
        ws[f"{col}{r}"].number_format = number_format


def priority_fill(ws, col: str, first: int, last: int) -> None:
    for r in range(first, last + 1):
        v = ws[f"{col}{r}"].value
        if v in PRIORITY_FILL:
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=PRIORITY_FILL[v])
            ws[f"{col}{r}"].font = Font(size=10, bold=True, color=INK)


def numbered(items: list[str]) -> str:
    return "\n".join(f"{i}. {t}" for i, t in enumerate(items, 1)) if items else ""


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--analysis", type=Path, required=True)
    ap.add_argument("--pages", type=Path)
    ap.add_argument("--gaps", type=Path)
    ap.add_argument("--suggestions", type=Path)
    ap.add_argument("-o", "--out", type=Path, required=True)
    ap.add_argument("--md", type=Path, help="Also write a short markdown summary")
    ap.add_argument("--site-name")
    args = ap.parse_args()

    an = json.loads(args.analysis.read_text())
    pages = json.loads(args.pages.read_text()) if args.pages and args.pages.exists() else {}
    gaps = json.loads(args.gaps.read_text()) if args.gaps and args.gaps.exists() else {}
    sug = json.loads(args.suggestions.read_text()) if args.suggestions and args.suggestions.exists() else {}
    psug = sug.get("pages", {})
    site_name = args.site_name or sug.get("site_name") or an["settings"].get("site") or (
        urlparse(an["pages_to_fetch"][0]).netloc if an.get("pages_to_fetch") else "Website")
    today = dt.date.today().isoformat()
    period = an.get("date_range") or "date range not stated (add --date-range to analyze_gsc.py)"

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_low = wb.create_sheet("Low CTR Pages")
    ws_p23 = wb.create_sheet("Page 2-3 Pages")
    ws_gap = wb.create_sheet("Keyword Gaps")
    ws_chk = wb.create_sheet("Content Chunks")
    ws_lnk = wb.create_sheet("Internal Links")
    ws_can = wb.create_sheet("Cannibalisation") if an.get("cannibalisation") else None
    ws_act = wb.create_sheet("Action Plan")
    ws_met = wb.create_sheet("Method & Notes")

    # Only pages that were fetched + gap-checked (analyze_gsc.py --top) go into the tabs
    focus = set(an.get("pages_to_fetch", []))
    low_list = [p for p in an["low_ctr_pages"] if p["page"] in focus]
    p23_list = [p for p in an["page23_pages"] if p["page"] in focus]
    omitted_low = len(an["low_ctr_pages"]) - len(low_list)
    omitted_p23 = len(an["page23_pages"]) - len(p23_list)
    if omitted_low or omitted_p23:
        an.setdefault("warnings", []).append(
            f"{omitted_low} more low-CTR page(s) and {omitted_p23} more page 2-3 page(s) qualified but were outside "
            f"--top and are not in this workbook. Re-run analyze_gsc.py with a higher --top to include them.")

    def as_list(v):
        return [v] if isinstance(v, str) else list(v or [])

    # ------------------------------------------------------------------ Method & Notes (first: other tabs reference it)
    s = an["settings"]
    ws_met["A1"] = "Method & Notes"
    ws_met["A1"].font = Font(size=16, bold=True, color=NAVY)
    meta_rows = [
        ["Site", site_name], ["GSC date range", period], ["Data mode", an["data_mode"]],
        ["Report generated", today], ["Benchmark", f"{s['benchmark']} CTR curve"],
        ["Low-CTR rule", f"Avg position < 11, impressions >= {s['min_impr']}, CTR < {s['ctr_threshold']:.0%} of expected CTR"],
        ["Page 2-3 rule", f"Avg position 11-30.9, impressions >= {s['min_impr_p23']}"],
        ["Min impressions per page+query row", s["min_pair_impr"]],
        ["Brand terms excluded", ", ".join(s["brand_terms"]) or "none given"],
        ["Expected CTR at position 5", s["expected_ctr_pos5"]],
    ]
    for i, (k, v) in enumerate(meta_rows, 3):
        ws_met.cell(row=i, column=1, value=k).font = Font(bold=True, size=10)
        ws_met.cell(row=i, column=2, value=v).alignment = WRAP
    pos5_ref = "'Method & Notes'!$B$12"
    ws_met["B12"].number_format = "0.00%"
    ws_met.column_dimensions["A"].width = 34
    ws_met.column_dimensions["B"].width = 90
    r = 15
    ws_met.cell(row=r, column=1, value="Expected CTR curve (position -> CTR)").font = Font(bold=True, color=NAVY)
    r += 1
    for k, v in s["ctr_curve"].items():
        ws_met.cell(row=r, column=1, value=f"Position {k}")
        ws_met.cell(row=r, column=2, value=v).number_format = "0.0%"
        ws_met.cell(row=r, column=2).alignment = Alignment(horizontal="left")
        r += 1
    r += 1
    notes = [
        ("How estimates work", "Extra clicks = impressions x expected CTR - current clicks. 'Clicks at position 5' uses "
         "today's impressions, so it is conservative: impressions usually grow when a page reaches page 1."),
        ("Average position caveat", "A page's average position blends all its queries. A page at 8.4 may rank #2 for one "
         "keyword and #25 for another - always read the query list."),
        ("Keyword status", "Used = phrase in body/headings and in title/H1/heading. Not in headings = only in body copy. "
         "Scattered terms = every word is there, never as the phrase. Partial = some words missing. Missing = most words missing."),
        ("Colour key", "Green cells = suggestions to review. Amber cells = for you to fill in (owner, status, notes)."),
        ("Re-measure", "Wait 4-6 weeks after changes, then compare the same date-range length in GSC (Compare mode)."),
    ]
    for k, v in notes + [("Warning", w) for w in an.get("warnings", [])]:
        ws_met.cell(row=r, column=1, value=k).font = Font(bold=True, size=10)
        ws_met.cell(row=r, column=2, value=v).alignment = WRAP
        r += 1
    r += 1
    ws_met.cell(row=r, column=1, value=ATTRIBUTION).font = Font(italic=True, color=MUTED, size=9)

    # ------------------------------------------------------------------ Low CTR Pages
    low_rows = []
    for p in low_list:
        url = p["page"]
        pg, g, sg = pages.get(url, {}), gaps.get(url, {}), psug.get(url, {})
        tq = an.get("page_queries", {}).get(url, [])[:6]
        titles, metas = as_list(sg.get("title_options")), as_list(sg.get("meta_options"))
        low_rows.append([
            url, p["clicks"], p["impressions"], None, p["position"], p["expected_ctr"], None, None, p["priority"],
            "\n".join(f"{q['query']} ({q['impressions']:,} impr, pos {q['position']})" for q in tq),
            pg.get("title", ""), None, pg.get("meta_description", ""), None,
            "\n".join(f"- {c}" for c in g.get("snippet_checks", [])),
            sg.get("diagnosis", ""), titles[0] if titles else "", None,
            titles[1] if len(titles) > 1 else "", None, metas[0] if metas else "", None,
            metas[1] if len(metas) > 1 else "", None,
            "\n".join(f"- {x}" for x in ([f"H1: {sg['h1_suggestion']}"] if sg.get("h1_suggestion") else []) + as_list(sg.get("other_fixes"))),
            "To do",
        ])
    hdr = [("Page URL", 42), ("Clicks", 9), ("Impressions", 12), ("CTR", 8), ("Avg position", 9),
           ("Expected CTR", 10), ("CTR vs expected", 10), ("Extra clicks at expected CTR", 12), ("Priority", 9),
           ("Top queries", 40), ("Current title", 38), ("Title chars", 7), ("Current meta description", 42),
           ("Meta chars", 7), ("Snippet issues found", 44), ("Diagnosis", 44), ("Suggested title (recommended)", 40),
           ("New title chars", 8), ("Alternative title", 36), ("Alt title chars", 8),
           ("Suggested meta description", 46), ("New meta chars", 8), ("Alternative meta description", 46),
           ("Alt meta chars", 8), ("Other fixes", 40), ("Status", 11)]
    last = table(ws_low, hdr, low_rows, suggest_cols=(16, 17, 19, 21, 23, 25), status_col=26)
    for rr in range(2, last + 1):
        ws_low[f"D{rr}"] = f"=IFERROR(B{rr}/C{rr},0)"
        ws_low[f"G{rr}"] = f"=IFERROR(D{rr}/F{rr},0)"
        ws_low[f"H{rr}"] = f"=MAX(0,ROUND(C{rr}*F{rr}-B{rr},0))"
        ws_low[f"L{rr}"] = f"=LEN(K{rr})"
        ws_low[f"N{rr}"] = f"=LEN(M{rr})"
        ws_low[f"R{rr}"] = f"=LEN(Q{rr})"
        ws_low[f"T{rr}"] = f"=LEN(S{rr})"
        ws_low[f"V{rr}"] = f"=LEN(U{rr})"
        ws_low[f"X{rr}"] = f"=LEN(W{rr})"
    if last >= 2:
        for col, f in (("C", "#,##0"), ("B", "#,##0"), ("D", "0.00%"), ("F", "0.00%"), ("G", "0%"), ("H", "#,##0"), ("E", "0.0")):
            fmt(ws_low, col, 2, last, f)
        priority_fill(ws_low, "I", 2, last)
        ws_low.conditional_formatting.add(f"G2:G{last}", ColorScaleRule(start_type="num", start_value=0, start_color="F87171",
                                                                         end_type="num", end_value=1, end_color="FFFFFF"))
        red = PatternFill("solid", fgColor="FECACA")
        for col, lim in (("L", 60), ("R", 60), ("T", 60), ("N", 160), ("V", 160), ("X", 160)):
            ws_low.conditional_formatting.add(f"{col}2:{col}{last}", CellIsRule(operator="greaterThan", formula=[str(lim)], fill=red))
        for col in ("N", "V", "X"):
            ws_low.conditional_formatting.add(f"{col}2:{col}{last}", FormulaRule(formula=[f"AND({col}2>0,{col}2<70)"], fill=red))

    # ------------------------------------------------------------------ Page 2-3 Pages
    p23_rows = []
    for p in p23_list:
        url = p["page"]
        pg, g, sg = pages.get(url, {}), gaps.get(url, {}), psug.get(url, {})
        t23 = as_list(sg.get("title_options"))
        q23 = [q for q in an.get("page_queries", {}).get(url, []) if 11 <= q["position"] < 31][:8]
        links = g.get("link_sources", [])
        p23_rows.append([
            url, None, p["clicks"], p["impressions"], p["position"], None, None, p["priority"],
            "\n".join(f"{q['query']} (pos {q['position']}, {q['impressions']:,} impr)" for q in q23),
            pg.get("word_count", ""), None,
            sg.get("diagnosis", ""), numbered(as_list(sg.get("actions"))),
            "\n".join(f"{short(l['source'])} -> anchor '{l['anchor']}'" for l in links),
            (t23[0] if t23 else "") + (f"\nH1: {sg['h1_suggestion']}" if sg.get("h1_suggestion") else ""), None,
            "To do",
        ])
    hdr = [("Page URL", 42), ("SERP page", 7), ("Clicks", 9), ("Impressions", 12), ("Avg position", 9), ("CTR", 8),
           ("Est. clicks at position 5", 12), ("Priority", 9), ("Keywords on page 2-3", 44), ("Word count", 9),
           ("Keyword gaps", 9), ("Diagnosis", 44), ("Push-to-page-1 plan", 60), ("Internal links to add", 44),
           ("Suggested title / H1 (optional)", 36), ("Title chars", 7), ("Status", 11)]
    last23 = table(ws_p23, hdr, p23_rows, suggest_cols=(12, 13, 14, 15), status_col=17)
    for rr in range(2, last23 + 1):
        ws_p23[f"B{rr}"] = f"=IF(E{rr}<21,2,3)"
        ws_p23[f"F{rr}"] = f"=IFERROR(C{rr}/D{rr},0)"
        ws_p23[f"G{rr}"] = f"=MAX(0,ROUND(D{rr}*{pos5_ref}-C{rr},0))"
        ws_p23[f"K{rr}"] = f"=COUNTIF('Keyword Gaps'!$A:$A,A{rr})"
        ws_p23[f"P{rr}"] = f'=IF(LEN(O{rr})=0,0,IFERROR(FIND(CHAR(10),O{rr})-1,LEN(O{rr})))'  # title line only
    if last23 >= 2:
        for col, f in (("C", "#,##0"), ("D", "#,##0"), ("E", "0.0"), ("F", "0.00%"), ("G", "#,##0")):
            fmt(ws_p23, col, 2, last23, f)
        priority_fill(ws_p23, "H", 2, last23)

    # ------------------------------------------------------------------ Keyword Gaps
    yn = lambda b: "Yes" if b else "No"  # noqa: E731
    gap_rows = []
    cluster_info = {}
    for url, e in gaps.items():
        for c in e.get("clusters", []):
            cluster_info[(url, c["id"])] = c
        for g in e.get("gaps", []):
            c = cluster_info.get((url, g.get("cluster")), {})
            gap_rows.append([url, g["query"], g["impressions"], g["clicks"], g["position"], g["status"], g["intent"],
                             yn(g["in_title"]), yn(g["in_h1"]), yn(g["in_headings"]), g["body_count"],
                             ", ".join(g["missing_terms"]), g.get("cluster", ""), c.get("chunk_type", ""),
                             c.get("placement", "")])
    hdr = [("Page URL", 42), ("Keyword (ranking query)", 34), ("Impressions", 11), ("Clicks", 8), ("Position", 8),
           ("Status", 14), ("Intent", 14), ("In title", 7), ("In H1", 7), ("In H2/H3", 8), ("Body mentions", 9),
           ("Missing words", 18), ("Chunk ID", 7), ("Fix type", 40), ("Where on the page", 40)]
    lastg = table(ws_gap, hdr, gap_rows)
    if lastg >= 2:
        fmt(ws_gap, "C", 2, lastg, "#,##0")
        status_colour = {"Missing": "FECACA", "Partial": "FED7AA", "Scattered terms": "FEF3C7", "Not in headings": "E0F2FE"}
        for rr in range(2, lastg + 1):
            v = ws_gap[f"F{rr}"].value
            if v in status_colour:
                ws_gap[f"F{rr}"].fill = PatternFill("solid", fgColor=status_colour[v])
            for col in "HIJ":
                if ws_gap[f"{col}{rr}"].value == "No":
                    ws_gap[f"{col}{rr}"].font = Font(size=10, color="B91C1C")

    # ------------------------------------------------------------------ Content Chunks
    written: dict[tuple, list] = {}
    for c in sug.get("chunks", []):
        written.setdefault((c.get("url"), c.get("cluster")), []).append(c)
    dups = [k for k, v in written.items() if len(v) > 1 and k[1]]
    if dups:
        print(f"  Note: {len(dups)} cluster(s) have more than one chunk - all are kept as separate rows.")
    chunk_rows, n = [], 0
    seen = set()

    def chunk_row(url, cid, w, c):
        return [url, cid, ", ".join(as_list(w.get("keywords")) or (c or {}).get("queries", [])),
                (c or {}).get("total_impressions"), w.get("chunk_type") or (c or {}).get("chunk_type", ""),
                w.get("placement") or (c or {}).get("placement", ""), w.get("heading", ""),
                w.get("content", "") or "(not written yet)", None, w.get("notes", ""), "To do"]

    for url, e in gaps.items():
        for c in e.get("clusters", []):
            seen.add((url, c["id"]))
            for w in written.get((url, c["id"]), [{}]):
                n += 1
                chunk_rows.append([n, *chunk_row(url, c["id"], w, c)])
    for (url, cid), ws_ in written.items():  # extra chunks Claude added outside the clusters
        if (url, cid) not in seen:
            for w in ws_:
                n += 1
                chunk_rows.append([n, *chunk_row(url, cid or "extra", w, None)])
    hdr = [("#", 5), ("Page URL", 38), ("Chunk ID", 7), ("Target keywords", 32), ("Keyword impressions", 11),
           ("Chunk type", 30), ("Where to add it", 34), ("Heading", 34), ("Content to add (ready to paste)", 80),
           ("Words", 7), ("Check before publishing", 32), ("Status", 11)]
    lastc = table(ws_chk, hdr, chunk_rows, suggest_cols=(8, 9), input_cols=(11,), status_col=12)
    for rr in range(2, lastc + 1):
        ws_chk[f"J{rr}"] = (f'=IF(OR(LEN(TRIM(I{rr}))=0,I{rr}="(not written yet)"),0,'
                            f'LEN(TRIM(I{rr}))-LEN(SUBSTITUTE(TRIM(I{rr})," ",""))+1)')
        fmt(ws_chk, "E", rr, rr, "#,##0")
        ws_chk.row_dimensions[rr].height = 150 if ws_chk[f"I{rr}"].value != "(not written yet)" else 30

    # ------------------------------------------------------------------ Internal Links
    link_rows = []
    for url, e in gaps.items():
        for l in e.get("link_sources", []):
            link_rows.append([url, l["source"], l["source_clicks"], l["anchor"], l["why"], "To do"])
    for url, sg in psug.items():
        for l in sg.get("internal_links", []):
            link_rows.append([url, l.get("source", ""), l.get("source_clicks", ""), l.get("anchor", ""),
                              l.get("why", "Suggested by review"), "To do"])
    hdr = [("Page to boost (link target)", 42), ("Add a link FROM this page", 42), ("Source page clicks", 10),
           ("Suggested anchor text", 30), ("Why this source", 50), ("Status", 11)]
    lastl = table(ws_lnk, hdr, link_rows, suggest_cols=(4,), status_col=6)
    fmt(ws_lnk, "C", 2, lastl, "#,##0")

    # ------------------------------------------------------------------ Cannibalisation
    if ws_can is not None:
        can_rows = []
        for c in an["cannibalisation"]:
            for i, p in enumerate(c["pages"]):
                can_rows.append([c["query"] if i == 0 else "", c["total_impressions"] if i == 0 else None,
                                 p["page"], p["impressions"], p["clicks"], p["position"],
                                 ("Pick ONE page to own this query. Point the others to it with an internal link "
                                  "using the query as anchor, and remove the query from their title/H1."
                                  if i == 0 else "")])
        hdr = [("Query", 32), ("Total impressions", 11), ("Competing page", 46), ("Impressions", 11), ("Clicks", 8),
               ("Position", 8), ("Suggested fix", 60)]
        lastk = table(ws_can, hdr, can_rows)
        fmt(ws_can, "B", 2, lastk, "#,##0")
        fmt(ws_can, "D", 2, lastk, "#,##0")

    # ------------------------------------------------------------------ Action Plan
    acts = []
    for p in low_list:
        acts.append([p["priority"], p["page"], "Rewrite title + meta description (see Low CTR Pages)", "CTR rewrite",
                     p["extra_clicks"]])
    for p in p23_list:
        e = gaps.get(p["page"], {})
        nc, nl = len(e.get("clusters", [])), len(e.get("link_sources", []))
        parts = []
        if nc:
            parts.append(f"Add {nc} content chunk{'s' if nc != 1 else ''}")
        if nl:
            parts.append(f"{nl} internal link{'s' if nl != 1 else ''}")
        what = (" + ".join(parts) + "; follow" if parts else "Follow") + " the push plan (see Page 2-3 Pages)"
        acts.append([p["priority"], p["page"], what, "Push to page 1", p["clicks_at_pos5"]])
    for p in low_list:
        nc = len(gaps.get(p["page"], {}).get("clusters", []))
        if nc:
            acts.append(["Medium", p["page"], f"Add {nc} content chunk{'s' if nc != 1 else ''} for keywords the page ranks for but doesn't use",
                         "Content chunks", None])
    for c in an.get("cannibalisation", [])[:10]:
        acts.append(["Medium", c["pages"][0]["page"], f"Fix cannibalisation for '{c['query']}' ({len(c['pages'])} pages competing)",
                     "Cannibalisation", None])
    rank = {"High": 0, "Medium": 1, "Low": 2}
    acts.sort(key=lambda a: (rank.get(a[0], 3), -(a[4] or 0)))
    act_rows = [[i, *a, "", "To do", "", ""] for i, a in enumerate(acts, 1)]
    hdr = [("#", 5), ("Priority", 9), ("Page", 44), ("Action", 60), ("Type", 15), ("Est. extra clicks / period", 12),
           ("Owner", 14), ("Status", 11), ("Date done", 11), ("Notes", 30)]
    lasta = table(ws_act, hdr, act_rows, input_cols=(7, 9, 10), status_col=8)
    if lasta >= 2:
        priority_fill(ws_act, "B", 2, lasta)
        fmt(ws_act, "F", 2, lasta, "#,##0")

    # ------------------------------------------------------------------ Summary
    ws = ws_sum
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    for col in "CDEF":
        ws.column_dimensions[col].width = 26
    ws["A1"] = f"Website Clicks Improver - {site_name}"
    ws["A1"].font = Font(size=18, bold=True, color=NAVY)
    ws["A2"] = f"GSC data: {period}  |  Mode: {an['data_mode']}  |  Generated {today}"
    ws["A2"].font = Font(size=10, color=MUTED)
    sm = an["summary"]
    kpis = [
        ("Clicks in the data", sm["total_clicks"], "#,##0"),
        ("Impressions in the data", sm["total_impressions"], "#,##0"),
        ("Site CTR", "=IFERROR(B4/B5,0)", "0.00%"),
        ("Low-CTR pages on page 1", "=COUNTA('Low CTR Pages'!A:A)-1", "0"),
        ("Extra clicks if they hit expected CTR", "=SUM('Low CTR Pages'!H:H)", "#,##0"),
        ("Pages stuck on Google page 2-3", "=COUNTA('Page 2-3 Pages'!A:A)-1", "0"),
        ("Clicks if those pages reach position 5", "=SUM('Page 2-3 Pages'!G:G)", "#,##0"),
        ("Total click opportunity (per period)", "=B8+B10", "#,##0"),
        ("Keyword gaps (ranking but not used)", "=COUNTA('Keyword Gaps'!B:B)-1", "0"),
        ("Content chunks written", "=COUNTIF('Content Chunks'!J:J,\">0\")", "0"),
        ("Actions done", "=COUNTIF('Action Plan'!H:H,\"Done\")&\" of \"&(COUNTA('Action Plan'!A:A)-1)", "@"),
    ]
    for i, (label, val, f) in enumerate(kpis, 4):
        a, b = ws.cell(row=i, column=1, value=label), ws.cell(row=i, column=2, value=val)
        a.fill = b.fill = KPI_FILL
        a.font = Font(size=11, color=INK)
        b.font = Font(size=12, bold=True, color=NAVY)
        b.number_format = f
        b.alignment = Alignment(horizontal="right")
        a.border = b.border = BORDER
    ws["A11"].font = ws["B11"].font = Font(size=12, bold=True, color="047857")
    r = 4 + len(kpis) + 1
    ws.cell(row=r, column=1, value="Top 10 actions").font = Font(size=13, bold=True, color=NAVY)
    top = [[a[2], a[0], short(a[1]), a[3], a[4]] for a in acts[:10]]
    hdr = [("Action", 46), ("Priority", 18), ("Page", 26), ("Type", 26), ("Est. extra clicks", 26)]
    for ci, (name, _) in enumerate(hdr, 1):
        c = ws.cell(row=r + 1, column=ci, value=name)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
    for ri, row in enumerate(top, r + 2):
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.alignment, c.border, c.font = WRAP, BORDER, Font(size=10)
        ws.cell(row=ri, column=5).number_format = "#,##0"
        if row[1] in PRIORITY_FILL:
            ws.cell(row=ri, column=2).fill = PatternFill("solid", fgColor=PRIORITY_FILL[row[1]])
    r = r + 2 + len(top) + 1
    ws.cell(row=r, column=1, value="How to use this workbook").font = Font(size=13, bold=True, color=NAVY)
    howto = [
        "1. Low CTR Pages: paste the suggested title + meta into your CMS (check the char counts stay green).",
        "2. Page 2-3 Pages: work through the push plan, starting with High priority rows.",
        "3. Content Chunks: add each chunk where it says, in your own voice. Check every fact flagged in 'Check before publishing'.",
        "4. Internal Links: add a contextual link from each source page using the suggested anchor.",
        "5. Action Plan: set an owner and mark Status as you go. Re-check GSC 4-6 weeks after each change.",
    ]
    for i, t in enumerate(howto, r + 1):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=False)
    r = r + len(howto) + 2
    for w in an.get("warnings", []):
        ws.cell(row=r, column=1, value=f"Note: {w}").font = Font(size=9, color="B45309")
        r += 1
    ws.cell(row=r + 1, column=1, value=ATTRIBUTION).font = Font(italic=True, size=9, color=MUTED)
    ws.sheet_view.showGridLines = False

    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor = {"Summary": NAVY, "Low CTR Pages": RED, "Page 2-3 Pages": AMBER,
                                           "Keyword Gaps": VIOLET, "Content Chunks": GREEN, "Internal Links": CYAN,
                                           "Action Plan": NAVY}.get(sheet.title, "9CA3AF")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"Saved workbook -> {args.out}")
    print(f"  Low CTR rows: {len(low_rows)} | Page 2-3 rows: {len(p23_rows)} | Gaps: {len(gap_rows)} | "
          f"Chunks: {len(chunk_rows)} ({sum(1 for c in chunk_rows if c[8] != '(not written yet)')} written) | "
          f"Links: {len(link_rows)} | Actions: {len(act_rows)}")
    missing = [p["page"] for p in an["low_ctr_pages"] if not psug.get(p["page"], {}).get("title_options")]
    if missing:
        print(f"  Missing title/meta suggestions for {len(missing)} low-CTR page(s).")

    if args.md:
        lines = [f"# Website Clicks Improver - {site_name}", "",
                 f"GSC data: {period} | generated {today}", "",
                 f"- **Low-CTR pages on page 1:** {sm['low_ctr_pages']} (~{sm['low_ctr_extra_clicks']:,} extra clicks at expected CTR)",
                 f"- **Pages on Google page 2-3:** {sm['page23_pages']} (~{sm['page23_clicks_at_pos5']:,} clicks if they reach position 5)",
                 f"- **Keyword gaps:** {len(gap_rows)} ranking queries not used properly on the page",
                 f"- **Content chunks:** {len(chunk_rows)}", "", "## Top actions", ""]
        for a in acts[:10]:
            est = f" (~{a[4]:,} clicks)" if a[4] else ""
            lines.append(f"1. **{a[0]}** - `{short(a[1])}` - {a[2]}{est}")
        lines += ["", f"_{ATTRIBUTION}_", ""]
        args.md.parent.mkdir(parents=True, exist_ok=True)
        args.md.write_text("\n".join(lines))
        print(f"Saved summary -> {args.md}")


if __name__ == "__main__":
    main()
